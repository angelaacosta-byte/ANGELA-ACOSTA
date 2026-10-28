"""Builds the cleaned Veritrade workbook (bobinas para conversion) with a full audit trail.
Usage: python build_clean.py <input.xlsx> <output.xlsx>"""
import sys, hashlib, unicodedata
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from limpieza_reglas import classify, MIN_DIAM, MIN_WIDTH

src, out = sys.argv[1], sys.argv[2]
md5 = hashlib.md5(open(src, 'rb').read()).hexdigest()

df = pd.read_excel(src, header=5)
df = df[df['Partida Aduanera'].notna()].copy()
cols = list(df.columns)
desc = ['Descripcion Comercial'] + [f'Descripcion{i}' for i in range(1, 6)]


def norm(s):
    return unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().upper()


txt = df[desc].fillna('').astype(str).agg(' | '.join, axis=1).map(norm)
res = txt.map(classify)
df['Decision'] = res.str[0]
df['Motivo'] = res.str[1]
uv = df['U$ FOB Tot'] / df['Kg Neto']
df['FOB por kg'] = uv.round(2)


def suggest(i):
    if df.at[i, 'Kg Neto'] >= 5000 and 0 < uv[i] <= 1.6:
        return 'Probable BOBINA (volumen alto y precio de bobina)'
    if df.at[i, 'Kg Neto'] < 1000 or uv[i] > 2.5:
        return 'Probable producto terminado (poco volumen o precio alto)'
    return 'Dudoso'


df['Decision'] = df['Decision'].map({'KEEP': 'CONSERVAR', 'REMOVE': 'ELIMINAR', 'REVIEW': 'REVISAR'})

# manual decisions by the domain expert (only applied to rows the rules sent to REVISAR)
import os
man_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'decisiones_manuales.csv')
n_manual = 0
if os.path.exists(man_path):
    man = pd.read_csv(man_path, dtype=str)
    key = lambda d: d['DUA / DAM'].astype(str) + '|' + d['Partida Aduanera'].astype(str).str[:10]
    dec = dict(zip(key(man), man['Decision'].str.upper()))
    k = key(df)
    for i in df.index[df.Decision == 'REVISAR']:
        d = dec.get(k[i])
        if d == 'SI':
            df.at[i, 'Decision'] = 'CONSERVAR'
            df.at[i, 'Motivo'] = 'revision manual de la experta: importador conversor de bobinas'
            n_manual += 1
        elif d == 'NO':
            df.at[i, 'Decision'] = 'ELIMINAR'
            df.at[i, 'Motivo'] = 'revision manual de la experta: no es bobina'
            n_manual += 1
keep = df[df.Decision == 'CONSERVAR']
rem = df[df.Decision == 'ELIMINAR']
rev = df[df.Decision == 'REVISAR'].copy()
rev['Sugerencia'] = [suggest(i) for i in rev.index]
rev['Tu decision (SI = es bobina / NO)'] = ''

HDR = PatternFill('solid', fgColor='1F4E78')
wb = Workbook()


def sheet(name, data, columns, color):
    ws = wb.create_sheet(name)
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=color)
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in data[columns].itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for j, col in enumerate(columns, 1):
        w = 60 if col in ('Descripcion Comercial', 'Motivo', 'Sugerencia') else 14
        ws.column_dimensions[get_column_letter(j)].width = w
        if col in ('Fecha', 'ETA', 'Fecha Embarque'):
            for c in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                for x in c:
                    x.number_format = 'yyyy-mm-dd'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


# summary first
ws = wb.active
ws.title = 'Resumen'
B = Font(bold=True)
rows = [
    ['Limpieza de datos Veritrade: bobinas de papel tissue para conversion', None],
    [None, None],
    ['Archivo de entrada', 'Partidas_4818200000-4803009000-4818900000.xlsx (version limpiada por Angela)'],
    ['Huella MD5 del archivo de entrada', md5],
    ['Periodo', f"{df['Fecha'].min():%Y-%m-%d} a {df['Fecha'].max():%Y-%m-%d}"],
    ['Registros de entrada', len(df)],
    ['Conservados (Bobinas_limpio)', len(keep)],
    ['Eliminados (Eliminados)', len(rem)],
    ['Por revisar (Revisar)', len(rev)],
    ['Decisiones manuales aplicadas (decisiones_manuales.csv)', n_manual],
    [None, None],
    ['Reglas, en orden de aplicacion', None],
    ['1', f'Diametro de rollo >= {MIN_DIAM} cm (D, RD, DIA, DIAMETRO, OD) -> BOBINA. Las bobinas miden ~100-130 cm; un rollo de consumo < 25 cm.'],
    ['2', 'Producto que no es papel para conversion (stickers, carton, papel de esterilizacion, panos, servilletas de fiesta, etc.) -> ELIMINAR, salvo senal fuerte de bobina.'],
    ['3', f'Ancho de rollo >= {MIN_WIDTH} cm (RW, ANCHO, WIDTH, W, FORMATO) o medida de 90-300 cm junto a gramaje -> BOBINA.'],
    ['4', 'Diametro < 40 cm -> ELIMINAR (rollo de consumo), salvo senal fuerte de bobina.'],
    ['5', 'Senal fuerte (JUMBO, ROLLO MADRE, PARA CONVERSION, FABRICACION, MATERIA PRIMA, toneladas) -> BOBINA aunque mencione empaque.'],
    ['6', 'Formato de consumo (hojas, paquetes, cajas, bolsas, interfoliado, rollos de 200 m, set, medidas de hoja) -> ELIMINAR.'],
    ['7', 'Otra senal de bobina (BOBINA, REEL, EN ROLLOS + GSM, USO INDUSTRIAL) sin empaque -> BOBINA.'],
    ['8', 'Uso domestico/comercial sin ninguna medida de bobina -> ELIMINAR.'],
    ['9', 'Lo demas -> REVISAR a mano (con sugerencia por volumen y precio por kg).'],
    ['10', 'Revision manual de la experta (decisiones_manuales.csv): las filas marcadas SI pasan a Bobinas_limpio; NO, a Eliminados.'],
    [None, None],
    ['Nota', 'Ninguna fila fue borrada sin motivo: cada fila eliminada aparece en la hoja Eliminados con la regla que la elimino.'],
]
for r in rows:
    ws.append(r)
ws['A1'].font = Font(bold=True, size=14)
ws['A12'].font = B
for r in range(3, 11):
    ws.cell(r, 1).font = B
ws.column_dimensions['A'].width = 48
ws.column_dimensions['B'].width = 120
ws.append([])
start = ws.max_row + 1
tab = pd.crosstab(df['Partida Aduanera'].astype('int64').astype(str), df['Decision'], margins=True, margins_name='Total')
ws.append(['Partida'] + list(tab.columns))
for c in ws[start]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = HDR
for idx, r in tab.iterrows():
    ws.append([idx] + [int(v) for v in r])

base = cols
sheet('Bobinas_limpio', keep, base + ['Motivo', 'FOB por kg'], '2E7D32')
sheet('Eliminados', rem, ['Motivo'] + base + ['FOB por kg'], 'C62828')
sheet('Revisar', rev.sort_values('Sugerencia'),
      ['Tu decision (SI = es bobina / NO)', 'Sugerencia', 'Motivo'] + base + ['FOB por kg'], 'EF6C00')
wb.save(out)
print(len(df), len(keep), len(rem), len(rev), md5)
print(rev['Sugerencia'].value_counts())
